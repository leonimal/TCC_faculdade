"""Gera o rascunho institucional do TCC a partir do conteúdo Markdown revisado.

O modelo original nunca é alterado. O arquivo gerado mantém as configurações de
página e os estilos do modelo, mas substitui completamente o texto herdado.
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "modelo" / "TCC_Modelo-UNIP2026Final.docx"
OUTPUT = ROOT / "entrega" / "TCC-IoT-UNIP-2026-revisado.docx"
SOURCES = [
    ROOT / "texto" / "01-introducao.md",
    ROOT / "texto" / "02-fundamentacao.md",
    ROOT / "texto" / "03-metodologia.md",
    ROOT / "texto" / "04-desenvolvimento.md",
    ROOT / "texto" / "05-resultados.md",
    ROOT / "texto" / "06-conclusao.md",
]

AUTHORS = ["ENRICO GARCIA SILVA COSTA", "LEONARDO PEDROZO ELIAS RODRIGUES"]
TITLE = (
    "DESENVOLVIMENTO E AVALIAÇÃO TÉCNICA DE UM SISTEMA IOT DE BAIXO "
    "CUSTO PARA MONITORAMENTO AMBIENTAL EM SALAS DE AULA"
)
CITY_YEAR = "ARAÇATUBA - SP\n2026"


def remove_template_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run_font(run, size=12, bold=None, italic=False) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    run.italic = italic


def set_paragraph_format(paragraph, align=None, space_before=0, space_after=0, line=1.5):
    if align is not None:
        paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line


def add_text(paragraph, text: str, *, size=12, bold=None, italic=False) -> None:
    """Add text with basic Markdown emphasis, retaining citations as plain text."""
    parts = re.split(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size, bold=True, italic=italic)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size, bold=bold, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size, bold=bold, italic=italic)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size, bold=bold, italic=italic)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    font = OxmlElement("w:rFonts")
    font.set(qn("w:ascii"), "Arial")
    font.set(qn("w:hAnsi"), "Arial")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    properties.extend((font, size))
    run.append(properties)
    field.append(run)
    paragraph._p.append(field)


def add_toc(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    paragraph._p.append(field)


def add_center(document: Document, text: str, *, size=12, bold=False, before=0, after=0):
    p = document.add_paragraph()
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.CENTER, before, after)
    for line_no, line in enumerate(text.split("\n")):
        if line_no:
            p.add_run().add_break()
        add_text(p, line, size=size, bold=bold)
    return p


def add_body(document: Document, text: str):
    p = document.add_paragraph()
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
    p.paragraph_format.first_line_indent = Cm(1.25)
    add_text(p, text)
    return p


def add_abstract_body(document: Document, source: Path) -> None:
    sections = [part.strip().replace("\n", " ") for part in source.read_text(encoding="utf-8").split("\n\n")]
    body, keywords = sections[1:3]
    p = add_body(document, body)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(12)
    p = document.add_paragraph()
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT, line=1.5)
    label, terms = keywords.split(":", 1)
    add_text(p, label + ":", bold=True)
    add_text(p, terms)


def configure_section(section) -> None:
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.header_distance = Cm(2)
    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    # O modelo contém um campo PAGE dentro de um controle de conteúdo no cabeçalho.
    # Limpar apenas o primeiro parágrafo não remove esse controle.
    for area in (section.header, section.footer):
        for child in list(area._element):
            area._element.remove(child)
        area.add_paragraph()


def start_page_count_at(section, number: int) -> None:
    properties = section._sectPr
    page_numbering = properties.find(qn("w:pgNumType"))
    if page_numbering is None:
        page_numbering = OxmlElement("w:pgNumType")
        properties.append(page_numbering)
    page_numbering.set(qn("w:start"), str(number))


def continue_page_count(section) -> None:
    page_numbering = section._sectPr.find(qn("w:pgNumType"))
    if page_numbering is not None:
        section._sectPr.remove(page_numbering)


def update_fields_on_open(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_heading(document: Document, text: str, level: int):
    p = document.add_paragraph(style=f"Heading {min(level, 3)}")
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT, space_before=12 if level == 1 else 6, space_after=6)
    p.paragraph_format.keep_with_next = True
    add_text(p, text, size=12)
    return p


def add_caption(document: Document, text: str):
    p = document.add_paragraph()
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=2, line=1.0)
    p.paragraph_format.keep_with_next = True
    add_text(p, text, size=10, bold=True)
    return p


def add_source(document: Document, text: str):
    p = document.add_paragraph()
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT, space_after=6, line=1.0)
    add_text(p, text, size=10)
    return p


def table_from_markdown(document: Document, rows: list[str]):
    matrix = []
    for row in rows:
        cells = [cell.strip() for cell in row.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        matrix.append(cells)
    if not matrix:
        return
    table = document.add_table(rows=0, cols=len(matrix[0]))
    table_properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    table_properties.append(borders)
    for row_index, cells in enumerate(matrix):
        row = table.add_row()
        row_properties = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_properties.append(cant_split)
        if row_index == 0:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            row_properties.append(header)
        row_cells = row.cells
        for col_index, cell_text in enumerate(cells):
            p = row_cells[col_index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_format(p, p.alignment, line=1.0)
            p.paragraph_format.keep_with_next = row_index < len(matrix) - 1
            add_text(p, cell_text, size=9, bold=row_index == 0)
    document.add_paragraph()


def render_markdown(document: Document, source: Path, *, start_new_page: bool):
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    first_heading = True
    while index < len(lines):
        raw = lines[index].strip()
        if not raw:
            index += 1
            continue
        if raw.startswith("# "):
            if start_new_page and first_heading:
                document.add_page_break()
            add_heading(document, raw[2:], 1)
            first_heading = False
        elif raw.startswith("## "):
            add_heading(document, raw[3:], 2)
        elif raw.startswith("### "):
            add_heading(document, raw[4:], 3)
        elif raw.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(lines[index].strip())
                index += 1
            table_from_markdown(document, rows)
            continue
        elif raw.startswith("**Fonte:") or raw.startswith("Fonte:"):
            add_source(document, raw.replace("**", ""))
        elif raw.startswith("**Tabela ") or raw.startswith("Tabela ") or raw.startswith("**Quadro "):
            add_caption(document, raw.replace("**", ""))
        elif raw.startswith("- "):
            p = document.add_paragraph()
            set_paragraph_format(p, WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2)
            add_text(p, "• " + raw[2:])
        elif re.match(r"^\d+\. ", raw):
            p = document.add_paragraph()
            set_paragraph_format(p, WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2)
            add_text(p, raw)
        else:
            # Internal project paths are useful for audit, not for the submitted text.
            raw = re.sub(r"`(?:experimento|software)/[^`]+`", "os artefatos digitais do projeto", raw)
            add_body(document, raw)
        index += 1


def render_references(document: Document):
    document.add_page_break()
    heading = add_heading(document, "REFERÊNCIAS", 1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text = (ROOT / "texto" / "07-referencias.md").read_text(encoding="utf-8")
    entries = [entry.replace("\n", " ").strip() for entry in re.split(r"\n\s*\n", text) if entry.strip()]
    for entry in entries[1:]:
        p = document.add_paragraph()
        set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT, space_after=12, line=1.0)
        p.paragraph_format.keep_together = True
        p.paragraph_format.left_indent = Cm(0)
        p.paragraph_format.first_line_indent = Cm(0)
        add_text(p, entry)


def main() -> None:
    OUTPUT.parent.mkdir(exist_ok=True)
    shutil.copy2(TEMPLATE, OUTPUT)
    document = Document(OUTPUT)
    remove_template_body(document)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    for level in (1, 2, 3):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    configure_section(document.sections[0])
    continue_page_count(document.sections[0])

    # Capa
    add_center(document, "UNIVERSIDADE PAULISTA - UNIP", bold=True, before=24)
    for author in AUTHORS:
        add_center(document, author, before=120 if author == AUTHORS[0] else 0)
    add_center(document, TITLE, bold=True, before=120)
    add_center(document, CITY_YEAR, before=150)
    pretext = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(pretext)
    start_page_count_at(pretext, 1)

    # Folha de rosto
    for author in AUTHORS:
        add_center(document, author, before=18 if author == AUTHORS[0] else 0)
    add_center(document, TITLE, bold=True, before=84)
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Cm(8)
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.0)
    add_text(p, "Trabalho de Conclusão de Curso para obtenção do título de Bacharel em Ciência da Computação apresentado à Universidade Paulista - UNIP.")
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Cm(8)
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT, line=1.0)
    add_text(p, "Orientador: Prof. Me. Rafael Marcelino de Jesus")
    add_center(document, CITY_YEAR, before=150)
    document.add_page_break()

    # Folha de aprovação: campos intencionalmente não preenchidos.
    add_center(document, "FOLHA DE APROVAÇÃO", bold=True, before=24)
    for author in AUTHORS:
        add_center(document, author, before=24 if author == AUTHORS[0] else 0)
    add_center(document, TITLE, bold=True, before=36)
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Cm(8)
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.0)
    add_text(p, "Trabalho de Conclusão de Curso para obtenção do título de Bacharel em Ciência da Computação apresentado à Universidade Paulista - UNIP.")
    add_center(document, "Aprovado em: ____/____/______", before=48)
    add_center(document, "BANCA EXAMINADORA", bold=True, before=24)
    add_center(document, "________________________________________\nProf.(a): ____________________________________\nUniversidade Paulista - UNIP", before=24)
    add_center(document, "________________________________________\nProf.(a): ____________________________________\nUniversidade Paulista - UNIP", before=18)
    document.add_page_break()

    add_center(document, "RESUMO", bold=True, before=24)
    add_abstract_body(document, ROOT / "texto" / "00-resumo.md")
    document.add_page_break()

    add_center(document, "ABSTRACT", bold=True, before=24)
    add_abstract_body(document, ROOT / "texto" / "00-abstract.md")
    document.add_page_break()

    add_center(document, "LISTA DE ABREVIATURAS E SIGLAS", bold=True, before=24)
    abbreviations = [
        ("API", "Application Programming Interface"),
        ("HTTP", "Hypertext Transfer Protocol"),
        ("I2C", "Inter-Integrated Circuit"),
        ("IoT", "Internet das Coisas"),
        ("JSON", "JavaScript Object Notation"),
        ("MAE", "Erro absoluto médio"),
        ("MAPE", "Erro percentual absoluto médio"),
        ("MySQL", "Sistema gerenciador de banco de dados relacional"),
        ("P95", "Percentil 95"),
        ("RMSE", "Raiz do erro quadrático médio"),
    ]
    for acronym, meaning in abbreviations:
        p = document.add_paragraph()
        set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT, space_after=2, line=1.0)
        add_text(p, f"{acronym} - {meaning}")
    document.add_page_break()

    add_center(document, "SUMÁRIO", bold=True, before=24)
    toc = document.add_paragraph()
    add_toc(toc)
    text_section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(text_section)
    continue_page_count(text_section)
    add_page_number(text_section.header.paragraphs[0])

    for index, source in enumerate(SOURCES):
        render_markdown(document, source, start_new_page=index > 0)
    render_references(document)

    update_fields_on_open(document)
    document.core_properties.title = TITLE.title()
    document.core_properties.author = "; ".join(title.title() for title in AUTHORS)
    document.core_properties.subject = "Trabalho de Conclusão de Curso - Ciência da Computação"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
